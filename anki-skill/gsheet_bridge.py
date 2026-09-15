#!/usr/bin/env python3
"""
Local bridge server for anki-manager.html's "Load from Google Sheets" picker.

anki-manager.html is a static file with no backend, so it can't do the
Desktop-app OAuth flow (google-auth-oauthlib's InstalledAppFlow +
run_local_server) that live_transcribe's Streamlit app uses for its Google
Doc import -- that flow needs a Python process to catch the OAuth redirect.
This script IS that process: it runs the Desktop-app sign-in once, caches
the token, and serves a tiny local JSON/CSV API that the page's JS fetches
from (http://localhost:8787 by default). No "Web application" OAuth client,
no Authorized JavaScript origin, and it works even opening anki-manager.html
directly via file://.

Setup (one-time):
    1. In Google Cloud Console (console.cloud.google.com), create an OAuth
       client of type "Desktop app" (APIs & Services > Credentials).
       Enable the Google Drive API for the project. OAuth consent screen
       can stay in "Testing" with your own account added as a test user.
    2. Download it and save as client_secret.json next to this file.
    3. pip install google-auth google-auth-oauthlib google-api-python-client openpyxl

Usage:
    python3 gsheet_bridge.py             # serves on http://localhost:8787
    python3 gsheet_bridge.py --port 9000  # then update GSHEET_BRIDGE_URL
                                           # in anki-manager.html to match

Leave this running in a terminal, then open anki-manager.html and use the
"Sign in with Google" / "Recent" controls in the Import CSV modal.
"""

import argparse
import csv
import io
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
RECENT_LIMIT = 10
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HERE = Path(__file__).parent
CLIENT_SECRETS_FILE = HERE / "client_secret.json"
TOKEN_FILE = HERE / "gsheet_token.json"  # cached credentials; gitignored, never commit

_workbook_cache: dict[str, object] = {}  # fileId -> openpyxl Workbook, cleared on sign-out


# ---------- Auth ----------

def load_cached_credentials():
    if not TOKEN_FILE.exists():
        return None
    info = json.loads(TOKEN_FILE.read_text())
    creds = Credentials.from_authorized_user_info(info, SCOPES)
    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())
        TOKEN_FILE.write_text(creds.to_json())
    return creds


def get_credentials():
    creds = load_cached_credentials()
    return creds if creds and creds.valid else None


def sign_in():
    if not CLIENT_SECRETS_FILE.exists():
        raise RuntimeError(
            f"Missing {CLIENT_SECRETS_FILE.name}. Create an OAuth client (type 'Desktop app') "
            "in Google Cloud Console, enable the Drive API, and download it as "
            f"{CLIENT_SECRETS_FILE.name} into {CLIENT_SECRETS_FILE.parent}."
        )
    flow = InstalledAppFlow.from_client_secrets_file(str(CLIENT_SECRETS_FILE), SCOPES)
    creds = flow.run_local_server(port=0)
    TOKEN_FILE.write_text(creds.to_json())
    return creds


def sign_out():
    _workbook_cache.clear()
    if TOKEN_FILE.exists():
        TOKEN_FILE.unlink()


# ---------- Drive ----------

def fetch_recent_sheets(creds) -> list[dict]:
    drive = build("drive", "v3", credentials=creds)
    resp = drive.files().list(
        q="mimeType='application/vnd.google-apps.spreadsheet' and trashed=false",
        orderBy="viewedByMeTime desc",
        pageSize=RECENT_LIMIT,
        fields="files(id, name)",
    ).execute()
    return resp.get("files", [])


def fetch_workbook(creds, file_id: str):
    """Download the whole spreadsheet as .xlsx (Drive's CSV export only
    gives the first tab; .xlsx gives every tab) and return it as an
    openpyxl Workbook, cached in memory for this server's lifetime."""
    if file_id in _workbook_cache:
        return _workbook_cache[file_id]
    drive = build("drive", "v3", credentials=creds)
    request = drive.files().export_media(fileId=file_id, mimeType=XLSX_MIME)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buf.seek(0)
    wb = load_workbook(buf, data_only=True, read_only=True)
    _workbook_cache[file_id] = wb
    return wb


def sheet_to_csv(workbook, sheet_name: str) -> str:
    ws = workbook[sheet_name]
    out = io.StringIO()
    writer = csv.writer(out)
    for row in ws.iter_rows(values_only=True):
        writer.writerow(["" if c is None else c for c in row])
    return out.getvalue()


# ---------- HTTP server ----------

class Handler(BaseHTTPRequestHandler):
    def _send_json(self, obj, status=200):
        body = json.dumps(obj).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _send_text(self, text, status=200, content_type="text/plain; charset=utf-8"):
        body = text.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _send_error(self, message, status=500):
        self._send_json({"error": message}, status=status)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        qs = parse_qs(parsed.query)
        try:
            if parsed.path == "/status":
                self._send_json({"signedIn": get_credentials() is not None})
            elif parsed.path == "/signin":
                sign_in()
                self._send_json({"signedIn": True})
            elif parsed.path == "/signout":
                sign_out()
                self._send_json({"signedIn": False})
            elif parsed.path == "/recent":
                creds = get_credentials()
                if not creds:
                    return self._send_error("Not signed in", 401)
                self._send_json({"files": fetch_recent_sheets(creds)})
            elif parsed.path == "/workbook":
                creds = get_credentials()
                if not creds:
                    return self._send_error("Not signed in", 401)
                file_id = qs.get("id", [""])[0]
                if not file_id:
                    return self._send_error("Missing ?id=", 400)
                wb = fetch_workbook(creds, file_id)
                self._send_json({"sheets": wb.sheetnames})
            elif parsed.path == "/workbook/csv":
                creds = get_credentials()
                if not creds:
                    return self._send_error("Not signed in", 401)
                file_id = qs.get("id", [""])[0]
                sheet = qs.get("sheet", [""])[0]
                if not file_id or not sheet:
                    return self._send_error("Missing ?id= or ?sheet=", 400)
                wb = fetch_workbook(creds, file_id)
                self._send_text(sheet_to_csv(wb, sheet), content_type="text/csv; charset=utf-8")
            else:
                self._send_error("Not found", 404)
        except Exception as e:  # noqa: BLE001 -- surface any failure to the page instead of a raw 500
            self._send_error(str(e), 500)

    def log_message(self, fmt, *args):
        print("[gsheet_bridge]", fmt % args)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--port", type=int, default=8787)
    args = parser.parse_args()

    server = ThreadingHTTPServer(("localhost", args.port), Handler)
    print(f"gsheet_bridge listening on http://localhost:{args.port}")
    if args.port != 8787:
        print("Note: update GSHEET_BRIDGE_URL near the top of the '// GOOGLE SHEETS PICKER' "
              "block in anki-manager.html to match this port.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
